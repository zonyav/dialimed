from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .. import groups
from ..models import COLUMNS, DASH, RunResult

FONT = "Bahnschrift Light"
FONT_HEAD = "Bahnschrift SemiBold"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=FONT_HEAD, color="FFFFFF", bold=False, size=11)
BODY_FONT = Font(name=FONT, size=11)
LINK_FONT = Font(name=FONT, size=11, color="1F6FEB", underline="single")

ALT_FILL = PatternFill("solid", fgColor="F4F6FA")
# заливка говорит одно: откуда взялось значение. Жёлтая — программа подобрала
# сама и это стоит перепроверить; красная — в контракте не было ничего, за что
# можно зацепиться, и вопрос не к программе, а к заказчику
GUESS_FILL = PatternFill("solid", fgColor="FFE1B3")
MISSING_FILL = PatternFill("solid", fgColor="F8CFCB")

HAIR = Side(style="hair", color="C9CEDA")
BORDER = Border(bottom=HAIR, left=HAIR, right=HAIR)

MONEY = '#,##0.00" ₽"'
INT = "#,##0"
PERCENT = '0.0" %"'
DATE = "DD.MM.YYYY"

MISSING_COLUMNS = ("№ РУ", "Производитель", "Держатель РУ в РФ",
                   "ИНН держателя РУ")
LINK_COLUMN = "Ссылка на ЕИС"
LINK_TEXT = "открыть в ЕИС"

# через месяц после прогона по цвету уже не вспомнить, что он означал,
# поэтому расшифровка висит примечанием на шапке «Производитель»
LEGEND_COLUMN = "Производитель"

# Сводка говорит «такие строки стоит просмотреть глазами», а найти их в файле
# было нечем: цвет тут не годится — номер РУ взят из самого контракта, ячейка
# по правилам белая. Поэтому примечание прямо в ячейке.
MISMATCH_SOURCE = "реестр РЗН (наименование не совпало)"
LEGEND = "\n".join((
    "Цвет ячейки:",
    "• без заливки — взято из реестра по номеру РУ, "
    "который указан в самом контракте;",
    "• жёлтая — программа подобрала сама (по обозначению из соседней строки "
    "или по срезу реестра); стоит перепроверить;",
    "• красная — в контракте нет ни номера РУ, ни модели, ни товарного "
    "знака: определять не по чему.",
))

LEFT = "left"
CENTER = "center"

# колонка -> (ширина, выравнивание, формат числа, переносить ли текст)
LAYOUT: dict[str, tuple[int, str, Optional[str], bool]] = {
    "Дата контракта": (13, CENTER, DATE, False),
    "Ссылка на ЕИС": (16, CENTER, None, False),
    "Наименование позиции": (42, LEFT, None, True),
    "Текст позиции из контракта": (58, LEFT, None, True),
    "НКМИ": (10, CENTER, None, False),
    "КТРУ": (23, CENTER, None, False),
    "№ РУ": (19, LEFT, None, False),
    "Производитель": (34, LEFT, None, True),
    "Держатель РУ в РФ": (30, LEFT, None, True),
    "ИНН держателя РУ": (16, CENTER, None, False),
    "Цена за ед., ₽": (17, CENTER, MONEY, False),
    "Количество": (12, CENTER, INT, False),
    "Сумма по позиции, ₽": (18, CENTER, MONEY, False),
    "Ставка НДС": (12, CENTER, None, False),
    "Цена контракта, ₽": (18, CENTER, MONEY, False),
    "НМЦК, ₽": (18, CENTER, MONEY, False),
    "Снижение, %": (26, LEFT, PERCENT, True),
    "Заказчик": (40, LEFT, None, True),
    "Поставщик": (34, LEFT, None, True),
    "ИНН поставщика": (16, CENTER, None, False),
    # характеристики не переносим: строка из десятка пар «имя: значение»
    # растягивала строку на пол-экрана, а рядом с ней всё остальное пусто
    "Характеристики": (70, LEFT, None, False),
}

DEFAULT_LAYOUT = (18, LEFT, None, False)


def build_workbook(result: RunResult, params_note: str = "") -> Workbook:
    # книга — только таблица позиций: сводка живёт на странице,
    # где её можно листать, сортировать и объединять производителей
    wb = Workbook()
    _sheet_positions(wb.active, result)
    # запрос виден в свойствах файла: через месяц по имени отчёта
    # уже не вспомнить, какие коды и период в нём лежат
    wb.properties.title = "Медизделия в контрактах ЕИС"
    wb.properties.creator = "Поиск медизделия"
    if params_note:
        wb.properties.description = params_note
    return wb


def _sheet_positions(ws: Worksheet, result: RunResult) -> None:
    ws.title = "Позиции"
    idx = {name: i for i, name in enumerate(COLUMNS, 1)}

    for name, i in idx.items():
        width, _align, _fmt, _wrap = LAYOUT.get(name, DEFAULT_LAYOUT)
        ws.column_dimensions[get_column_letter(i)].width = width
        cell = ws.cell(row=1, column=i, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal=CENTER, vertical="center",
                                   wrap_text=True)
        if name == LEGEND_COLUMN:
            note = Comment(LEGEND, "Поиск медизделия")
            note.width, note.height = 400, 150
            cell.comment = note
    ws.row_dimensions[1].height = 38

    # в таблицу идёт объединённое имя производителя — то же, что
    # в списке на странице
    titles = groups.assign(r.pos.manufacturer for r in result.rows)

    # раскладку колонок считаем один раз: иначе на отчёт в тысячу строк
    # приходится двадцать тысяч одинаковых поисков по словарю
    plan = [(name, i, *LAYOUT.get(name, DEFAULT_LAYOUT)[1:]) for name, i in idx.items()]

    n = 0
    for row in result.rows:
        n += 1
        r = n + 1
        data = row.as_dict()
        striped = n % 2 == 0
        for name, i, align, fmt, wrap in plan:
            value = data[name]
            if name == "Производитель":
                value = titles.get(row.pos.manufacturer, value)
            cell = ws.cell(row=r, column=i, value=value)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(horizontal=align, vertical="top",
                                       wrap_text=wrap)
            if fmt and isinstance(value, (int, float)):
                cell.number_format = fmt
            elif fmt == DATE and value not in (None, DASH):
                cell.number_format = fmt
            if striped:
                cell.fill = ALT_FILL

        clue = row.pos.has_clue
        for name in MISSING_COLUMNS:
            cell = ws.cell(row=r, column=idx[name])
            if not cell.value or cell.value == DASH:
                cell.fill = GUESS_FILL if clue else MISSING_FILL
            elif name == "Производитель" and not row.pos.from_contract_number:
                cell.fill = GUESS_FILL

        if row.pos.manufacturer_source == MISMATCH_SOURCE and row.pos.ru_registry_name:
            cell = ws.cell(row=r, column=idx["Производитель"])
            note = Comment("\n".join((
                "Реестр называет это изделие иначе:",
                row.pos.ru_registry_name[:300],
                "",
                "Производитель взят по номеру РУ из контракта. Стоит открыть "
                "контракт и проверить, тот ли это номер.",
            )), "Поиск медизделия")
            note.width, note.height = 380, 150
            cell.comment = note

        card = row.pos.rzn_url
        if card:
            ru = ws.cell(row=r, column=idx["№ РУ"])
            ru.hyperlink = card
            ru.font = LINK_FONT

        link = ws.cell(row=r, column=idx[LINK_COLUMN])
        if isinstance(link.value, str) and link.value.startswith("http"):
            link.hyperlink = link.value
            link.value = LINK_TEXT
            link.font = LINK_FONT

    if n:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{n + 1}"
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False


def report_name(codes: list[str]) -> str:
    """Имя файла отчёта: код первой позиции и время прогона."""

    tag = codes[0].replace(".", "_") if codes else "папка"
    # запрос бывает не кодом, а словами («рускан 70п») — в имени файла им
    # нельзя ни пробелов, ни того, что Windows не пустит в имя
    tag = re.sub(r"[^0-9A-Za-zА-Яа-яЁё_-]+", "_", tag).strip("_") or "папка"
    return f"медизделия_{tag[:60]}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"


def save_report(result: RunResult, path: str | Path, params_note: str = "") -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = build_workbook(result, params_note)
    wb.save(path)
    return path
